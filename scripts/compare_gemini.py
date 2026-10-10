# -*- coding: utf-8 -*-
"""
compare_gemini.py — يقارن حلول Gemini العمياء بمفاتيح البنك وتصنيفه، ويكتب:
  • فحص_Gemini/مقارنة_Gemini.csv و ملخص_فحص_Gemini.md (للمراجعين)
  • أعمدة «الفحص المسبق المستقل» في ورقة «مراجعة البنود» ومؤشراتها في «لوحة المتابعة» داخل متتبع_الاعتماد.xlsx
لا يكتب في أعمدة قرارات المراجعين البشر.
"""
import csv
import json
import pathlib
import re
import sys
from collections import Counter

from openpyxl import load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
import build_bank as bb  # noqa: E402
from gemini_review import card_list, load_all  # noqa: E402

ROOT = pathlib.Path(__file__).resolve().parents[1]
KIT = ROOT / "05_الاعتماد"
OUT = KIT / "02_التدقيق_العلمي_الرياضي" / "فحص_Gemini"
TRACKER = KIT / "متتبع_الاعتماد.xlsx"
AGREE, DISAGREE, HUMAN, MISSING = "اتفاق", "اختلاف", "يحتاج حكماً بشرياً", "لا نتيجة"
PRI = {1: "1 اختلاف في الحل", 2: "2 اتفاق مع ملاحظة", 3: "3 حكم بشري", 4: "4 اتفاق دون ملاحظات"}
NUM_RE = re.compile(r"[-−]?\d+(?:\.\d+)?(?:\s*/\s*\d+(?:\.\d+)?)?")


def to_num(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    m = NUM_RE.search(str(v).replace(",", ""))
    if not m:
        return None
    s = m.group(0).replace("−", "-").replace(" ", "")
    if "/" in s:
        a, b = s.split("/")
        return float(a) / float(b) if float(b) else None
    return float(s)


def close(a, b, tol):
    if a is None or b is None:
        return False
    t = float(tol) if tol else max(1e-6, 1e-6 * abs(b))
    return abs(a - b) <= t + 1e-9


def key_text(it):
    f = it["format"]
    if f in ("mc", "select_explain"):
        return str(it.get("key"))
    if f == "numeric":
        return str(it["answer"]["value"])
    if f == "multi_numeric":
        return json.dumps([x["value"] for x in it.get("fields", [])], ensure_ascii=False)
    if f in ("dragdrop_order", "dragdrop_categorize", "dragdrop_slots", "grid_plot"):
        return json.dumps(it.get("key"), ensure_ascii=False)
    return "مفتوح (دليل تصحيح)"


def judge(it, g):
    """يعيد (الحالة، ملاحظة قصيرة)."""
    f = it["format"]
    if g is None:
        return MISSING, ""
    ans = g.get("answer")
    if f == "cr_explain":
        return HUMAN, ""
    if f in ("mc", "select_explain"):
        letter = str(ans).strip().upper()[:1] if ans is not None else ""
        return (AGREE if letter == str(it.get("key")).upper() else DISAGREE), ""
    if f == "numeric":
        a = it["answer"]
        targets = [a["value"]] + list(a.get("accept") or [])
        gv = to_num(ans)
        return (AGREE if any(close(gv, to_num(t), a.get("tol")) for t in targets) else DISAGREE), ""
    if f == "multi_numeric":
        vals = ans if isinstance(ans, list) else [ans]
        fl = it.get("fields", [])
        ok = len(vals) == len(fl) and all(close(to_num(v), to_num(x["value"]), x.get("tol")) for v, x in zip(vals, fl))
        return (AGREE if ok else DISAGREE), ""
    cards = card_list(it.get("cards"))
    by_text = {txt: cid for cid, txt in cards}
    ids = {cid for cid, _ in cards}
    to_id = lambda ref: ref if ref in ids else by_text.get(str(ref), str(ref))
    if f == "dragdrop_order":
        key = [to_id(k) for k in it.get("key") or []]
        got = [to_id(x) for x in ans] if isinstance(ans, list) else []
        return (AGREE if got == key else DISAGREE), ""
    if f == "dragdrop_categorize":
        norm = lambda s: re.sub(r"[\s$]", "", str(s))   # Gemini قد يحذف علامات $ من أسماء الفئات
        key = {to_id(k): norm(v) for k, v in (it.get("key") or {}).items()}
        got = {to_id(k): norm(v) for k, v in ans.items()} if isinstance(ans, dict) else {}
        diff = [k for k in key if got.get(k) != key[k]]
        return (AGREE if not diff else DISAGREE), (f"بطاقات مختلفة: {', '.join(diff)}" if diff else "")
    if f == "dragdrop_slots":
        key = {str(k): to_id(v) for k, v in (it.get("key") or {}).items()}
        got = {str(k): to_id(v) for k, v in ans.items()} if isinstance(ans, dict) else {}
        if got == key:
            return AGREE, ""
        expr = it.get("check") or it.get("verify")
        if expr and got:
            byid = {c["id"]: c for c in bb.norm_cards(it.get("cards") or [])}
            vals = {f"s{s}": (byid[cid]["value"] if cid in byid else None) for s, cid in got.items()}
            try:
                code = bb.js_to_py(expr) if it.get("check") else expr
                if bool(bb.run_verify(code, vals)):
                    return AGREE, "إجابة بديلة صحيحة يقبلها شرط التصحيح"
            except Exception:
                pass
        return DISAGREE, ""
    if f == "grid_plot":
        key = sorted(tuple(p) for p in it.get("key") or [])
        try:
            got = sorted(tuple(float(c) for c in p) for p in ans)
        except Exception:
            return DISAGREE, "صيغة نقاط غير مقروءة"
        if [tuple(map(float, p)) for p in key] == got:
            return AGREE, ""
        if it.get("verify"):
            try:
                if bool(bb.run_verify(it["verify"], {"P": [tuple(int(c) if float(c).is_integer() else c for c in p) for p in got]})):
                    return AGREE, "نقاط بديلة يقبلها شرط التحقق"
            except Exception:
                pass
        return DISAGREE, ""
    return HUMAN, ""


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    res = json.loads((OUT / "results.json").read_text(encoding="utf-8"))
    rows = []
    for it, ctx in load_all():
        g = res.get(it["id"])
        st, note = judge(it, g)
        topic_ok = g is not None and str(g.get("topic", "")).strip().upper() == it["topic"]
        cog_ok = g is not None and str(g.get("cognitive", "")).strip().lower() == it["cognitive"]
        cls = "—" if g is None else "اتفاق تام" if topic_ok and cog_ok else "اتفاق في الموضوع" if topic_ok else "اتفاق في المستوى" if cog_ok else "اختلاف"
        issues = [] if g is None else [*g.get("issues", []), *g.get("language_issues", []), *g.get("sensitivity_issues", [])]
        conf = None if g is None else g.get("confidence")
        unamb = None if g is None else g.get("unambiguous")
        flag = bool(g and (g.get("issues") or unamb is False or (conf is not None and conf < 0.7)))
        pri = 1 if st == DISAGREE else 3 if st in (HUMAN, MISSING) else 2 if flag or issues else 4
        src = "PSI" if ctx else ("تجريبي" if "-X" in it["id"] else "تدريب")
        rows.append({
            "id": it["id"], "source": src, "format": it["format"], "topic": it["topic"], "cognitive": it["cognitive"],
            "key": key_text(it), "gemini": "" if g is None else json.dumps(g.get("answer"), ensure_ascii=False),
            "status": st, "note": note, "confidence": conf, "priority": PRI[pri],
            "g_topic": "" if g is None else g.get("topic", ""), "g_cog": "" if g is None else g.get("cognitive", ""), "classification": cls,
            "math_issues": " · ".join(g.get("issues", [])) if g else "", "language_issues": " · ".join(g.get("language_issues", [])) if g else "",
            "sensitivity_issues": " · ".join(g.get("sensitivity_issues", [])) if g else "",
        })
    # CSV للمراجعين
    with open(OUT / "مقارنة_Gemini.csv", "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)
    c_st = Counter(r["status"] for r in rows)
    c_pri = Counter(r["priority"] for r in rows)
    c_cls = Counter(r["classification"] for r in rows)
    n = len(rows)
    auto = [r for r in rows if r["status"] in (AGREE, DISAGREE)]
    topic_agree = sum(1 for r in rows if r["classification"] in ("اتفاق تام", "اتفاق في الموضوع"))
    cog_agree = sum(1 for r in rows if r["classification"] in ("اتفاق تام", "اتفاق في المستوى"))
    dis = [r for r in rows if r["status"] == DISAGREE]
    flagged = [r for r in rows if r["priority"] == PRI[2]]
    md = [f"# ملخص الفحص المسبق المستقل بنموذج Gemini 3.1 Pro\n",
          "حلّ نموذج من عائلة أخرى (Gemini عبر Antigravity) كل بند دون رؤية المفتاح أو الحل، وصنّفه، ورصد ملاحظاته. هذه النتائج ترتيب لأولويات المراجعة البشرية، لا قرار اعتماد.\n",
          "## الأرقام",
          f"- البنود: {n} · لها نتيجة: {n - c_st.get(MISSING, 0)}",
          f"- البنود القابلة للمقارنة الآلية: {len(auto)} · اتفاق: {c_st.get(AGREE, 0)} · اختلاف: {c_st.get(DISAGREE, 0)} · نسبة الاتفاق: {c_st.get(AGREE, 0) / max(1, len(auto)):.1%}",
          f"- تحتاج حكماً بشرياً (أسئلة مفتوحة): {c_st.get(HUMAN, 0)}",
          f"- اتفاق تصنيف الموضوع: {topic_agree / n:.1%} · اتفاق المستوى المعرفي: {cog_agree / n:.1%}",
          "- أولويات المراجعة: " + " · ".join(f"{k}: {c_pri.get(k, 0)}" for k in PRI.values()),
          "\n## الأولوية 1: اختلاف الحل عن المفتاح\n", "| البند | المصدر | النوع | المفتاح | جواب Gemini | ملاحظة Gemini |", "|---|---|---|---|---|---|"]
    for r in dis:
        md.append(f"| {r['id']} | {r['source']} | {r['format']} | {r['key'][:60]} | {r['gemini'][:60]} | {(r['math_issues'] or r['note'])[:120]} |")
    md += ["\n## الأولوية 2: اتفاق مع ملاحظة\n", "| البند | ملاحظة رياضية | ملاحظة لغوية أو ثقافية |", "|---|---|---|"]
    for r in flagged:
        md.append(f"| {r['id']} | {r['math_issues'][:140]} | {(r['language_issues'] + ' ' + r['sensitivity_issues']).strip()[:140]} |")
    (OUT / "ملخص_فحص_Gemini.md").write_text("\n".join(md) + "\n", encoding="utf-8")

    # المتتبع
    wb = load_workbook(TRACKER)
    ws = wb["مراجعة البنود"]
    heads = ["فحص Gemini: الحل", "جواب Gemini", "أولوية المراجعة", "تصنيف Gemini", "اتفاق التصنيف", "ملاحظات Gemini"]
    thin = Side(style="thin", color="C9CFDA")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for k, h in enumerate(heads, 16):
        c = ws.cell(row=1, column=k, value=h)
        c.font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
        c.fill = PatternFill("solid", fgColor="2A7F8F")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
    for col, wdt in zip("PQRSTU", (14, 22, 20, 16, 16, 50)):
        ws.column_dimensions[col].width = wdt
    by_id = {r["id"]: r for r in rows}
    calc = PatternFill("solid", fgColor="E3F1F3")
    last = ws.max_row
    for rr in range(2, last + 1):
        r = by_id.get(ws.cell(row=rr, column=1).value)
        if not r:
            continue
        vals = [r["status"], r["gemini"][:200], r["priority"], f"{r['g_topic']} / {r['g_cog']}", r["classification"],
                " · ".join(x for x in (r["math_issues"], r["note"], r["language_issues"], r["sensitivity_issues"]) if x)[:500]]
        for k, v in enumerate(vals, 16):
            c = ws.cell(row=rr, column=k, value=v)
            c.font = Font(name="Arial", size=10)
            c.fill = calc
            c.border = border
            c.alignment = Alignment(horizontal="right", vertical="top", wrap_text=False)
    ws.auto_filter.ref = f"A1:U{last}"
    ws.conditional_formatting.add(f"P2:P{last}", CellIsRule(operator="equal", formula=['"اختلاف"'], fill=PatternFill("solid", fgColor="F8D3D8")))
    ws.conditional_formatting.add(f"P2:P{last}", CellIsRule(operator="equal", formula=['"اتفاق"'], fill=PatternFill("solid", fgColor="D9F0DF")))
    # لوحة المتابعة
    d = wb["لوحة المتابعة"]
    title = "الفحص المسبق المستقل (Gemini 3.1 Pro)"
    found = [r for r in range(1, d.max_row + 1) if d.cell(row=r, column=1).value == title]
    base = found[0] if found else d.max_row + 2          # إعادة التشغيل تحدّث الكتلة نفسها ولا تكررها
    d.cell(row=base, column=1, value=title).font = Font(name="Arial", bold=True, size=12, color="8A1538")
    rng = f"'مراجعة البنود'!$P$2:$P${last}"
    pr = f"'مراجعة البنود'!$R$2:$R${last}"
    cl = f"'مراجعة البنود'!$T$2:$T${last}"
    items_ = [("اتفاق في الحل", f'=COUNTIF({rng},"{AGREE}")'), ("اختلاف في الحل", f'=COUNTIF({rng},"{DISAGREE}")'),
              ("يحتاج حكماً بشرياً", f'=COUNTIF({rng},"{HUMAN}")'),
              ("نسبة الاتفاق في الحل", f"=IF(B{base + 1}+B{base + 2}=0,0,B{base + 1}/(B{base + 1}+B{base + 2}))"),
              ("اتفاق تصنيف الموضوع", f'=(COUNTIF({cl},"اتفاق تام")+COUNTIF({cl},"اتفاق في الموضوع"))/COUNTA({cl})'),
              ("اتفاق المستوى المعرفي", f'=(COUNTIF({cl},"اتفاق تام")+COUNTIF({cl},"اتفاق في المستوى"))/COUNTA({cl})')]
    items_ += [(f"أولوية {v}", f'=COUNTIF({pr},"{v}")') for v in PRI.values()]
    for k, (lab, f) in enumerate(items_, base + 1):
        d.cell(row=k, column=1, value=lab).font = Font(name="Arial", size=10)
        c = d.cell(row=k, column=2, value=f)
        c.font = Font(name="Arial", size=10)
        c.fill = PatternFill("solid", fgColor="EEF1F6")
        c.border = border
        if "نسبة" in lab or "اتفاق تصنيف" in lab or "اتفاق المستوى" in lab:
            c.number_format = "0.0%"
    wb.save(TRACKER)
    print(f"بنود {n} | اتفاق {c_st.get(AGREE, 0)} | اختلاف {c_st.get(DISAGREE, 0)} | حكم بشري {c_st.get(HUMAN, 0)} | بلا نتيجة {c_st.get(MISSING, 0)}")
    print(f"اتفاق الموضوع {topic_agree / n:.1%} | اتفاق المستوى {cog_agree / n:.1%} | أولوية 2 (ملاحظات): {len(flagged)}")


if __name__ == "__main__":
    main()
