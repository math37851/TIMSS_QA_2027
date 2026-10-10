# -*- coding: utf-8 -*-
"""
tracker_validations.py — يعيد القوائم المنسدلة في متتبع_الاعتماد.xlsx عبر أسماء معرّفة (L_STATUS، L_SCI…).

لماذا؟ Excel يخزّن القائمة التي تشير مباشرة إلى ورقة أخرى بامتداد (x14) لا تدعمه openpyxl فتحذفه عند الحفظ.
القائمة التي تشير إلى اسم معرّف تُخزَّن بالصيغة القياسية فتبقى. شغّله بعد أي تعديل للمتتبع بـ openpyxl،
ثم أعد الحساب في Excel.
"""
import pathlib
import sys
import warnings

from openpyxl import load_workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

TRACKER = pathlib.Path(__file__).resolve().parents[1] / "05_الاعتماد" / "متتبع_الاعتماد.xlsx"
LISTS = "القوائم"
# الاسم ← (عمود القائمة في ورقة القوائم، [(الورقة، عمود الخلايا)])
SPEC = {
    "L_STATUS": ("A", [("المهام", "K")]),
    "L_SCI": ("B", [("مراجعة البنود", "J")]),
    "L_LANG": ("C", [("مراجعة البنود", "L")]),
    "L_CURR": ("D", [("مراجعة البنود", "N")]),
    "L_SEV": ("E", [("سجل الملاحظات", "F")]),
    "L_ISTAT": ("F", [("سجل الملاحظات", "H")]),
    "L_TRACK": ("G", [("سجل الملاحظات", "C")]),
    "L_DRES": ("H", [("اختبار الأجهزة", "D")]),
}
EXTRA_ROWS = {"المهام": 60}


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    warnings.simplefilter("ignore")
    wb = load_workbook(TRACKER)
    wl = wb[LISTS]
    for name, (col, targets) in SPEC.items():
        last = max(r for r in range(2, wl.max_row + 1) if wl[f"{col}{r}"].value not in (None, "")) if wl[f"{col}2"].value else 2
        ref = f"'{LISTS}'!${col}$2:${col}${last}"
        if name in wb.defined_names:
            del wb.defined_names[name]
        wb.defined_names[name] = DefinedName(name, attr_text=ref)
        for sheet, tcol in targets:
            ws = wb[sheet]
            ws.data_validations.dataValidation = [dv for dv in ws.data_validations.dataValidation if tcol not in str(dv.sqref)]
            dv = DataValidation(type="list", formula1=name, allow_blank=True, showErrorMessage=True,
                                errorTitle="قيمة غير صالحة", error="اختر قيمة من القائمة.")
            ws.add_data_validation(dv)
            dv.add(f"{tcol}2:{tcol}{ws.max_row + EXTRA_ROWS.get(sheet, 0)}")
    wb.save(TRACKER)
    for ws in wb.worksheets:
        n = len(ws.data_validations.dataValidation)
        if n:
            print(ws.title, "| قوائم منسدلة:", n)


if __name__ == "__main__":
    main()
